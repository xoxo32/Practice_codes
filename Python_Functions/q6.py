# Q6: Study groups - split the Excel sheet into files, or find your group

import os
from openpyxl import load_workbook

# The Excel file must be in the same folder as this script
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(SCRIPT_DIR, "CSFB-FT11 Study Groups.xlsx")
OUTPUT_FOLDER = r"C:\Users\User\Desktop\Moringa-School-Labs\Python_Functions\study_groups"


def name_from_email(email):
    """niven.egaire@student... -> Niven Egaire (dots become spaces, digits removed)"""
    local_part = email.split("@")[0]
    local_part = local_part.replace(".", " ")
    local_part = "".join(ch for ch in local_part if not ch.isdigit())
    return local_part.strip().title()


def read_groups():
    """Read the spreadsheet into a dictionary {group_number: [emails]}"""
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    groups = {}
    current_group = None

    for row in sheet.iter_rows(min_row=3, values_only=True):  # rows 1-2 are titles/headers
        group_cell, email_cell = row[0], row[1]

        # The group label ("Grp 1") only appears on the first row of each group
        if group_cell:
            current_group = str(group_cell).split()[-1]
            groups[current_group] = []

        # Skip blank separator rows
        if email_cell and current_group:
            groups[current_group].append(email_cell.strip())

    return groups


def split_groups(groups):
    """Write one text file per group"""
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)

    for group_number, emails in groups.items():
        filename = os.path.join(OUTPUT_FOLDER, f"group{group_number}.txt")

        with open(filename, "w") as file:
            for email in emails:
                file.write(f"{name_from_email(email)} - {email}\n")

        print(f"Created group{group_number}.txt with {len(emails)} members")


def find_group(groups):
    """Search for a student by email or name and show their whole group"""
    query = input("Enter your email or name: ").strip().lower()

    if not query:
        print("You didn't type anything.")
        return

    found = False

    for group_number, emails in groups.items():
        for email in emails:
            name = name_from_email(email)

            if "@" in query:
                # Full email typed: compare against the whole email address
                is_match = query == email.lower()
            else:
                # Name typed: match the start of the full name or of any word in it
                # e.g. "sc" matches "Scott Mwangi" but not "Oscar Otieno"
                is_match = (" " + name.lower()).find(" " + query) != -1

            if is_match:
                found = True
                print(f"\n{name} ({email}) is in Group {group_number}.")
                print(f"Members of Group {group_number}:")
                for member_email in emails:
                    print(f"  - {name_from_email(member_email)} - {member_email}")

    if not found:
        print("No student found with that name or email. Check the spelling and try again.")


def main():
    groups = read_groups()

    while True:
        print("\nSelect which action to take place:")
        print("1. Split the groups")
        print("2. Find your group via email or name")
        print("3. Exit")

        choice = input("Enter 1, 2 or 3: ").strip()

        if choice == "1":
            split_groups(groups)
        elif choice == "2":
            find_group(groups)
        elif choice == "3":
            print("Goodbye!")
            break
        else:
            print("Invalid choice. Please enter 1, 2 or 3.")


if __name__ == "__main__":
    main()